import sqlite3
import openpyxl

EXCEL_FILE = "Gastos mes a mes 2026.xlsx"
DB_FILE = "financas.db"

MESES = [
    "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
    "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"
]

def safe_float(val):
    try:
        return float(val) if val is not None else 0.0
    except (ValueError, TypeError):
        return 0.0

conn = sqlite3.connect(DB_FILE)
cursor = conn.cursor()

# Criação das tabelas
cursor.execute("""
CREATE TABLE IF NOT EXISTS despesas (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    categoria TEXT NOT NULL,
    mes TEXT NOT NULL,
    ano INTEGER NOT NULL,
    valor REAL NOT NULL,
    UNIQUE(categoria, mes, ano)
)
""")

cursor.execute("""
CREATE TABLE IF NOT EXISTS rendas (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    mes TEXT NOT NULL,
    ano INTEGER NOT NULL,
    valor REAL NOT NULL,
    UNIQUE(mes, ano)
)
""")

wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
ws = wb["Mensal"]

# 1. Migração das Despesas (linhas 4 a 23)
for r in range(4, 24):
    cat = ws.cell(row=r, column=1).value
    if cat and str(cat).strip():
        categoria = str(cat).strip()
        for idx, mes in enumerate(MESES, start=2):
            val = safe_float(ws.cell(row=r, column=idx).value)
            cursor.execute("""
                INSERT OR REPLACE INTO despesas (categoria, mes, ano, valor)
                VALUES (?, ?, 2026, ?)
            """, (categoria, mes, val))

# 2. Migração da Renda (linha 27)
for idx, mes in enumerate(MESES, start=2):
    renda_val = safe_float(ws.cell(row=27, column=idx).value)
    cursor.execute("""
        INSERT OR REPLACE INTO rendas (mes, ano, valor)
        VALUES (?, 2026, ?)
    """, (mes, renda_val))

conn.commit()
conn.close()
print("Migração concluída com sucesso! Banco 'financas.db' gerado.")